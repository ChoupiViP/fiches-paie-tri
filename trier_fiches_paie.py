#!/usr/bin/env python3
"""
Tri de fiches de paie (PDF numériques) + récapitulatif annuel.

Tout reste en local. Les originaux ne sont jamais modifiés : le script COPIE
les fichiers dans un dossier de sortie, rangés par année.

Utilisation en ligne de commande :
    pip install -r requirements.txt
    python trier_fiches_paie.py "C:\\chemin\\vers\\les\\pdf" -o resultat
    python trier_fiches_paie.py "C:\\chemin\\vers\\les\\pdf" -o resultat --debug

Pour une version avec fenêtre : python app.py

Gérer une mise en page inconnue :
    1. lancer avec --debug : le texte brut de chaque PDF est écrit dans
       resultat/debug/*.txt
    2. repérer comment le net imposable / la période / l'employeur y apparaissent
    3. ajouter une expression dans un fichier patterns.json (à côté du
       programme), sans toucher au code. Exemple :
       {
         "net_imposable": ["montant fiscal"],
         "employeur": ["agence\\s*:\\s*(.+)"]
       }
"""
import argparse
import copy
import hashlib
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:  # pragma: no cover
    sys.exit("Il manque des modules. Installe-les avec :\n  pip install -r requirements.txt")

# --------------------------------------------------------------------------
# Motifs de lecture (plusieurs variantes par champ pour couvrir plusieurs
# mises en page). Les motifs sont testés dans l'ordre, du plus fiable au moins
# fiable.
# --------------------------------------------------------------------------
PATTERNS_DEFAUT = {
    # libellés (regex, insensible à la casse) suivis d'un montant
    "net_imposable": [
        r"net\s+imposable",
        r"net\s+fiscal",
        r"revenu\s+net\s+imposable",
        r"salaire\s+net\s+imposable",
        r"net\s+soumis\s+[àa]\s+l.?imp[oô]t",
    ],
    "cumul_net_imposable": [
        r"cumul\w*\s+.{0,25}net\s+imposable",
        r"net\s+imposable\s+(?:cumul\w*|annuel|depuis)",
        r"net\s+imposable\s+cumul",
        r"cumul\s+annuel\s+net\s+imposable",
    ],
    "brut": [
        r"salaire\s+brut",
        r"total\s+brut",
        r"brut\s+total",
        r"^\s*brut\b",
        r"r[ée]mun[ée]ration\s+brute",
    ],
    "net_a_payer": [
        r"net\s+[àa]\s+payer\s+avant\s+imp[oô]t",
        r"net\s+[àa]\s+payer",
        r"net\s+pay[ée]",
        r"net\s+vers[ée]",
    ],
    # employeur : motifs avec un groupe capturant le nom
    "employeur": [
        r"(?:employeur|raison\s+sociale|soci[ée]t[ée]|entreprise|agence)\s*[:\-]\s*(.+)",
    ],
}
PATTERNS = copy.deepcopy(PATTERNS_DEFAUT)

# Agences / enseignes courantes d'intérim (reconnues n'importe où dans le texte)
EMPLOYEURS_CONNUS = [
    "adecco", "randstad", "manpower", "synergie", "crit", "partnaire", "supplay",
    "proman", "daco", "start people", "temporis", "interaction", "samsic",
    "actual", "gi group", "walters", "terre d'int", "triangle", "domino",
    "inter-contrats", "groupe daco", "ctt", "adéquat", "adequat",
]

MOIS = {
    "janvier": 1, "fevrier": 2, "février": 2, "mars": 3, "avril": 4, "mai": 5,
    "juin": 6, "juillet": 7, "aout": 8, "août": 8, "septembre": 9,
    "octobre": 10, "novembre": 11, "decembre": 12, "décembre": 12,
}
NOMS_MOIS = ["", "janv", "fév", "mars", "avr", "mai", "juin", "juil", "août",
             "sept", "oct", "nov", "déc"]

SEP_MILLIERS = r"[ \u00a0\u202f.]"
MONTANT = re.compile(
    r"(?<![\d/])(-?\d{1,3}(?:" + SEP_MILLIERS + r"\d{3})*[,.]\d{2}|-?\d+[,.]\d{2})(?![\d/])"
)

_perso_charge = False


def dossier_programme():
    """Dossier où se trouve le programme (ou le .exe une fois compilé)."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def charger_patterns_perso(log=print):
    """Ajoute les motifs de patterns.json (s'il existe) en tête de liste."""
    global _perso_charge
    if _perso_charge:
        return
    p = dossier_programme() / "patterns.json"
    if p.exists():
        try:
            extra = json.loads(p.read_text(encoding="utf-8"))
            for cle, liste in extra.items():
                PATTERNS.setdefault(cle, [])
                PATTERNS[cle] = list(liste) + PATTERNS[cle]
            log(f"Motifs personnalisés chargés depuis {p.name}")
        except Exception as e:
            log(f"Attention : patterns.json illisible ({e})")
    _perso_charge = True


def en_nombre(txt):
    s = re.sub(r"[ \u00a0\u202f]", "", txt)
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def lire_texte(pdf_path):
    morceaux = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            morceaux.append(page.extract_text() or "")
    return "\n".join(morceaux)


def montants_ligne(ligne):
    return [en_nombre(m) for m in MONTANT.findall(ligne)]


def chercher_montants(lignes, motifs, exclure=None, apres_label_seulement=True):
    """Retourne la liste des montants de la 1re ligne qui correspond à un motif
    (ou de la ligne suivante si la ligne du libellé n'en contient pas)."""
    for motif in motifs:
        rx = re.compile(motif, re.I)
        for i, ligne in enumerate(lignes):
            if exclure and re.search(exclure, ligne, re.I):
                continue
            m = rx.search(ligne)
            if not m:
                continue
            reste = ligne[m.end():] if apres_label_seulement else ligne
            vals = montants_ligne(reste)
            if not vals and i + 1 < len(lignes):
                vals = montants_ligne(lignes[i + 1])
            if vals:
                return vals
    return None


def detecter_periode(texte):
    """Retourne (annee, mois) ou None. Plusieurs façons de lire la période."""
    t = texte.lower()
    # 1) "du 01/03/2026 au 31/03/2026" -> on prend la date de fin
    m = re.search(
        r"du\s+(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})\s+au\s+(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})", t)
    if m:
        mois, an = int(m.group(5)), int(m.group(6))
        an += 2000 if an < 100 else 0
        if 1 <= mois <= 12:
            return an, mois
    # 2) "période : 03/2026" / "bulletin de paie 03/2026" / "mois : 03-2026"
    m = re.search(r"(?:p[ée]riode|mois|paie|salaire|bulletin)[^\n\d]{0,30}(\d{1,2})[/.\-](\d{4})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(2)), int(m.group(1))
    # 3) "mars 2026" (nom du mois), 1re occurrence
    noms = "|".join(sorted(MOIS, key=len, reverse=True))
    m = re.search(r"\b(" + noms + r")\s+(?:de\s+)?(\d{4})\b", t)
    if m:
        return int(m.group(2)), MOIS[m.group(1)]
    # 4) mm/aaaa isolé
    m = re.search(r"\b(0?[1-9]|1[0-2])[/.\-](20\d{2})\b", t)
    if m:
        return int(m.group(2)), int(m.group(1))
    return None


def nettoyer_nom(nom):
    nom = re.sub(r"\s{2,}.*", "", nom.strip())          # coupe à la 1re grosse coupure
    nom = re.sub(r"\b(siret|siren|ape|naf|adresse|tel|urssaf).*", "", nom, flags=re.I)
    nom = re.sub(r"[^\w\s&'\-]", " ", nom)
    nom = re.sub(r"\s+", " ", nom).strip(" -_")
    return nom[:40]


def detecter_employeur(texte, lignes):
    bas = texte.lower()
    # 1) libellé explicite "Employeur : XXX"
    for motif in PATTERNS["employeur"]:
        m = re.search(motif, texte, re.I)
        if m:
            nom = nettoyer_nom(m.group(1))
            if len(nom) >= 2:
                return nom
    # 2) enseignes connues
    for nom in EMPLOYEURS_CONNUS:
        if nom in bas:
            return nom.title()
    # 3) 1re ligne "utile" du haut du document (souvent le nom de l'entreprise)
    for ligne in lignes[:8]:
        l = ligne.strip()
        if len(l) < 3 or re.search(r"bulletin|paie|salaire|p[ée]riode|n°|\d{5}", l, re.I):
            continue
        if re.search(r"[A-Za-zÀ-ÿ]{3,}", l):
            nom = nettoyer_nom(l)
            if len(nom) >= 3:
                return nom
    return None


def hash_fichier(p):
    h = hashlib.sha1()
    with open(p, "rb") as f:
        for bloc in iter(lambda: f.read(65536), b""):
            h.update(bloc)
    return h.hexdigest()


def nom_sur(nom):
    return re.sub(r"[^\w\-]+", "_", nom).strip("_") or "Inconnu"


def analyser(pdf_path, dossier_debug=None):
    texte = lire_texte(pdf_path)
    if dossier_debug is not None:
        (dossier_debug / (pdf_path.stem + ".txt")).write_text(texte, encoding="utf-8")
    res = {"fichier": pdf_path, "avert": []}
    if len(texte.strip()) < 40:
        res["avert"].append("pas de texte (scan ? il faudrait de l'OCR)")
        return res
    lignes = [l for l in texte.splitlines() if l.strip()]

    per = detecter_periode(texte)
    res["annee"], res["mois"] = per if per else (None, None)
    if not per:
        res["avert"].append("période non trouvée")

    res["employeur"] = detecter_employeur(texte, lignes)
    if not res["employeur"]:
        res["avert"].append("employeur non trouvé")

    # net imposable du mois : on ignore les lignes de cumul
    v = chercher_montants(lignes, PATTERNS["net_imposable"], exclure=r"cumul|annuel")
    res["net_imposable"] = v[0] if v else None
    if not v:
        res["avert"].append("net imposable non trouvé")
    elif len(v) > 1:
        res["avert"].append(f"plusieurs montants sur la ligne ({', '.join(f'{x:.2f}' for x in v)}) : 1er retenu")

    vc = chercher_montants(lignes, PATTERNS["cumul_net_imposable"])
    res["cumul"] = vc[-1] if vc else None

    vb = chercher_montants(lignes, PATTERNS["brut"])
    res["brut"] = vb[0] if vb else None
    vn = chercher_montants(lignes, PATTERNS["net_a_payer"])
    res["net_a_payer"] = vn[0] if vn else None
    return res


def traiter(dossier, sortie="tri_fiches", debug=False, log=print):
    """Trie les PDF de `dossier`, écrit le résultat dans `sortie`.

    Retourne un dict : xlsx (Path), sortie (Path), nb_pdf, doublons, lues,
    a_verifier, totaux (année -> montant).
    Lève ValueError si le dossier est introuvable ou sans PDF.
    """
    charger_patterns_perso(log)
    src = Path(dossier)
    out = Path(sortie)
    if not src.is_dir():
        raise ValueError(f"Dossier introuvable : {src}")
    out.mkdir(parents=True, exist_ok=True)
    dbg = None
    if debug:
        dbg = out / "debug"
        dbg.mkdir(exist_ok=True)

    out_res = out.resolve()
    pdfs = sorted(
        p for p in src.rglob("*")
        if p.is_file() and p.suffix.lower() == ".pdf" and out_res not in p.resolve().parents
    )
    if not pdfs:
        raise ValueError("Aucun PDF trouvé dans ce dossier.")
    log(f"{len(pdfs)} PDF trouvés, lecture en cours…")

    vus = {}
    resultats, doublons = [], []
    for p in pdfs:
        h = hash_fichier(p)
        if h in vus:
            doublons.append((p, vus[h]))
            continue
        vus[h] = p
        try:
            resultats.append(analyser(p, dbg))
        except Exception as e:
            resultats.append({"fichier": p, "avert": [f"erreur de lecture : {e}"]})

    # ----- rangement (copie) -----
    utilises = set()
    for r in resultats:
        ok = r.get("annee") and r.get("mois")
        if ok:
            dest_dir = out / str(r["annee"])
            base = f"{r['annee']}-{r['mois']:02d}_{nom_sur(r.get('employeur') or 'Inconnu')}"
        else:
            dest_dir = out / "A_verifier"
            base = r["fichier"].stem
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / f"{base}.pdf"
        n = 2
        while dest in utilises or dest.exists():
            dest = dest_dir / f"{base}_{n}.pdf"
            n += 1
        utilises.add(dest)
        shutil.copy2(r["fichier"], dest)
        r["classe_dans"] = dest

    # ----- agrégats -----
    par_cle = defaultdict(list)
    for r in resultats:
        if r.get("annee"):
            par_cle[(r["annee"], r.get("employeur") or "Inconnu")].append(r)

    wb = Workbook()
    gras = Font(bold=True)

    ws = wb.active
    ws.title = "Récap annuel"
    ws.append(["Année", "Employeur", "Nb fiches", "Net imposable (somme des fiches)",
               "Cumul lu sur la dernière fiche", "Écart", "Remarque"])
    totaux_annee = defaultdict(float)
    for (an, emp) in sorted(par_cle):
        fiches = par_cle[(an, emp)]
        somme = sum(f["net_imposable"] or 0 for f in fiches)
        manquants = sum(1 for f in fiches if f["net_imposable"] is None)
        cumuls = [(f["mois"], f["cumul"]) for f in fiches if f.get("cumul") is not None]
        cumul = max(cumuls)[1] if cumuls else None
        ecart = round(cumul - somme, 2) if cumul is not None else None
        rem = []
        if manquants:
            rem.append(f"{manquants} fiche(s) sans net imposable lu")
        if ecart not in (None, 0) and abs(ecart) > 1:
            rem.append("écart avec le cumul : fiches manquantes ou lecture à vérifier")
        ws.append([an, emp, len(fiches), round(somme, 2), cumul, ecart, " ; ".join(rem)])
        totaux_annee[an] += somme
    ws.append([])
    ws.append(["TOTAL PAR ANNÉE"])
    ws.cell(ws.max_row, 1).font = gras
    for an in sorted(totaux_annee):
        ws.append([an, "Tous employeurs", None, round(totaux_annee[an], 2)])

    wd = wb.create_sheet("Détail")
    wd.append(["Fichier d'origine", "Classé dans", "Année", "Mois", "Employeur", "Brut",
               "Net imposable", "Net à payer", "Cumul net imposable", "Avertissements"])
    for r in sorted(resultats, key=lambda x: (x.get("annee") or 0, x.get("mois") or 0)):
        wd.append([r["fichier"].name, str(r["classe_dans"].relative_to(out)), r.get("annee"),
                   r.get("mois"), r.get("employeur"), r.get("brut"), r.get("net_imposable"),
                   r.get("net_a_payer"), r.get("cumul"), " ; ".join(r["avert"])])

    wm = wb.create_sheet("Mois manquants")
    wm.append(["Employeur", "Année", "Mois sans fiche (entre la 1re et la dernière fiche)"])
    par_emp = defaultdict(set)
    for (an, emp), fiches in par_cle.items():
        for f in fiches:
            par_emp[emp].add((f["annee"], f["mois"]))
    for emp in sorted(par_emp):
        mois_vus = sorted(par_emp[emp])
        a0, m0 = mois_vus[0]
        a1, m1 = mois_vus[-1]
        a, m = a0, m0
        manques = defaultdict(list)
        while (a, m) <= (a1, m1):
            if (a, m) not in par_emp[emp]:
                manques[a].append(NOMS_MOIS[m])
            m += 1
            if m == 13:
                a, m = a + 1, 1
        for an in sorted(manques):
            wm.append([emp, an, ", ".join(manques[an])])
    wm.append([])
    wm.append(["Remarque : seuls les mois situés entre la première et la dernière fiche d'un employeur "
               "sont signalés. Pour de l'intérim à la semaine, plusieurs fiches par mois sont normales."])

    wv = wb.create_sheet("À vérifier")
    wv.append(["Fichier", "Problème"])
    nb_verif = 0
    for r in resultats:
        if r["avert"]:
            wv.append([r["fichier"].name, " ; ".join(r["avert"])])
            nb_verif += 1
    for p, orig in doublons:
        wv.append([p.name, f"doublon exact de {orig.name} (ignoré)"])

    for feuille in wb.worksheets:
        for c in feuille[1]:
            c.font = gras
        for col in feuille.columns:
            largeur = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            feuille.column_dimensions[col[0].column_letter].width = min(largeur + 2, 60)

    chemin_xlsx = out / "recapitulatif_fiches_paie.xlsx"
    wb.save(chemin_xlsx)

    nb_ok = sum(1 for r in resultats if r.get("annee") and r.get("net_imposable") is not None)
    log(f"{len(pdfs)} PDF trouvés, {len(doublons)} doublon(s) ignoré(s), "
        f"{nb_ok}/{len(resultats)} fiches lues complètement.")
    for an in sorted(totaux_annee):
        log(f"  {an} : net imposable cumulé des fiches = {totaux_annee[an]:,.2f} €".replace(",", " "))
    log(f"Récapitulatif : {chemin_xlsx}")
    log("Ouvre l'onglet « À vérifier » pour les fiches mal lues.")
    log("IMPORTANT : compare ces totaux avec le cumul de la dernière fiche de l'année "
        "avant de les reporter sur la déclaration.")

    return {
        "xlsx": chemin_xlsx, "sortie": out, "nb_pdf": len(pdfs), "doublons": len(doublons),
        "lues": nb_ok, "a_verifier": nb_verif + len(doublons), "totaux": dict(totaux_annee),
    }


def main():
    ap = argparse.ArgumentParser(description="Trie des fiches de paie PDF et fait un récap annuel.")
    ap.add_argument("dossier", help="dossier contenant les PDF (sous-dossiers inclus)")
    ap.add_argument("-o", "--sortie", default="tri_fiches", help="dossier de sortie")
    ap.add_argument("--debug", action="store_true", help="écrit le texte brut extrait de chaque PDF")
    args = ap.parse_args()
    try:
        traiter(args.dossier, args.sortie, args.debug)
    except ValueError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()
