import pytest

import trier_fiches_paie as c


def test_en_nombre():
    assert c.en_nombre("1 234,56") == 1234.56
    assert c.en_nombre("1.234,56") == 1234.56
    assert c.en_nombre("1234.56") == 1234.56
    assert c.en_nombre("-12,50") == -12.5


def test_arrondi_euro():
    assert c.arrondi_euro(1432.18) == 1432
    assert c.arrondi_euro(1432.5) == 1433
    assert c.arrondi_euro(2902.18) == 2902
    assert c.arrondi_euro(0) == 0


def test_cases_declarant():
    assert c.CASES_DECLARANT["1"][1] == "1AJ"
    assert c.CASES_DECLARANT["2"][1] == "1BJ"
    assert c.CASES_DECLARANT["3"][1] == "1CJ"
    assert c.CASES_DECLARANT["4"][1] == "1DJ"


def test_autres_cases_declarant():
    assert c.cases_pour("1") == {"libelle": "Déclarant 1", "salaires": "1AJ", "hs": "1GH", "pas": "8HV"}
    assert c.cases_pour("2")["hs"] == "1HH" and c.cases_pour("2")["pas"] == "8IV"
    assert c.cases_pour("3")["hs"] == "1IH" and c.cases_pour("3")["pas"] == "8JV"
    assert c.cases_pour("4")["hs"] == "1JH" and c.cases_pour("4")["pas"] == "8KV"


def test_rappels_suivent_la_personne():
    texte1 = " ".join(c.rappels(c.cases_pour("1")))
    assert "1AK" in texte1 and "8HV" in texte1 and "1GH" in texte1
    texte2 = " ".join(c.rappels(c.cases_pour("2")))
    assert "1BK" in texte2 and "8IV" in texte2 and "1HH" in texte2


def test_montants_ligne():
    assert c.montants_ligne("Net imposable 1 432,18") == [1432.18]
    assert c.montants_ligne("Net imposable 1 432,18 15 000,00") == [1432.18, 15000.0]
    assert c.montants_ligne("Période du 01/03/2026 au 31/03/2026") == []


@pytest.mark.parametrize("texte, attendu", [
    ("Période du 01/03/2026 au 31/03/2026", (2026, 3)),
    ("Période : 07/2024", (2024, 7)),
    ("Bulletin de salaire - mois de décembre 2025", (2025, 12)),
    ("Bulletin de paie mars 2023", (2023, 3)),
    ("rien d'utile ici", None),
])
def test_periode(texte, attendu):
    assert c.detecter_periode(texte) == attendu


def test_employeur():
    texte = "Employeur : SARL Dupont Informatique\nautre ligne"
    assert c.detecter_employeur(texte, texte.splitlines()) == "SARL Dupont Informatique"
    texte = "ADECCO France\nBulletin de paie"
    assert c.detecter_employeur(texte, texte.splitlines()) == "Adecco"


def _pdf(chemin, lignes):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    cv = canvas.Canvas(str(chemin), pagesize=A4)
    y = 800
    for ligne in lignes:
        cv.drawString(40, y, ligne)
        y -= 18
    cv.save()


def _jeu_de_pdf(src):
    src.mkdir()
    # trois mises en page différentes
    _pdf(src / "a1.pdf", ["ADECCO France", "BULLETIN DE PAIE", "Période du 01/03/2026 au 31/03/2026",
                          "Salaire brut 1 850,00", "Net imposable 1 432,18"])
    _pdf(src / "a2.pdf", ["ADECCO France", "BULLETIN DE PAIE", "Période du 01/05/2026 au 31/05/2026",
                          "Salaire brut 1 900,00", "Net imposable 1 470,00",
                          "Net imposable cumulé 2 902,18"])
    _pdf(src / "b1.pdf", ["Employeur : SARL Dupont Informatique", "Bulletin de salaire - mois de octobre 2025",
                          "Total brut 2 100,50", "Net fiscal 1 650,25"])
    _pdf(src / "c1.pdf", ["Société Martin Frères", "Période : 07/2024", "Net imposable", "1 210,00"])
    # doublon exact
    (src / "a1_copie.pdf").write_bytes((src / "a1.pdf").read_bytes())


def test_bout_en_bout(tmp_path):
    pytest.importorskip("reportlab")
    src = tmp_path / "pdf"
    _jeu_de_pdf(src)

    res = c.traiter(src, tmp_path / "out", log=lambda m: None)

    assert res["nb_pdf"] == 5
    assert res["doublons"] == 1
    assert res["lues"] == 4
    assert res["case"] == "1AJ"
    assert res["totaux"][2026] == pytest.approx(2902.18)
    assert res["totaux"][2025] == pytest.approx(1650.25)
    assert res["totaux"][2024] == pytest.approx(1210.0)
    assert (tmp_path / "out" / "2026" / "2026-03_Adecco.pdf").exists()
    assert res["xlsx"].exists()
    assert res["html"].exists()
    # l'onglet d'aide existe, en premier, et contient les montants arrondis
    from openpyxl import load_workbook
    wb = load_workbook(res["xlsx"])
    assert wb.sheetnames[0] == "Aide déclaration"
    texte = [str(v) for row in wb["Aide déclaration"].iter_rows(values_only=True) for v in row if v is not None]
    assert "2902" in texte and "1650" in texte and "1210" in texte
    # 2026 : mois d'avril sans fiche chez Adecco -> à vérifier
    assert res["aide"][2026]["raisons"]
    # 2024 : une fiche lue, aucun problème
    assert res["aide"][2024]["raisons"] == []


def test_declarant_conjoint(tmp_path):
    pytest.importorskip("reportlab")
    src = tmp_path / "pdf"
    _jeu_de_pdf(src)
    res = c.traiter(src, tmp_path / "out", log=lambda m: None, declarant="2")
    assert res["case"] == "1BJ"
    assert "1BJ" in res["html"].read_text(encoding="utf-8")


def test_employeur_ignore_les_montants():
    # régression : « employeur -14,62 » ne doit pas être pris pour un nom d'employeur
    texte = "JOB & BOX\nAutres contributions dues par l'employeur -14,62\nTotal brut 273,28"
    assert c.detecter_employeur(texte, texte.splitlines()) == "Job & Box"


def _pdf_tableau_cumuls(chemin, debut="28/09/2026", fin="30/09/2026", net=("222,95", "1 456,78"),
                        impot=("18,50", "18,50"), hs=("12,00", "12,00")):
    """Fiche d'intérim fictive avec un tableau « Cumuls » : titres sur une ligne,
    montants alignés à droite sous chaque titre, lignes Mensuel / Annuel.
    net, impot, hs : (valeur du mois, valeur annuelle) écrites dans les colonnes correspondantes."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    cv = canvas.Canvas(str(chemin), pagesize=A4)
    cv.setFont("Helvetica", 9)
    haut = [
        (40, "JOB & BOX"),
        (40, "BULLETIN DE PAYE"),
        (40, f"Période de paye du {debut} au {fin}"),
        (40, "Autres contributions dues par l'employeur -14,62"),
        (40, "Total Brut total * 273,28"),
        (40, "Total NET A PAYER AVANT IMPOTS SUR LE REVENU 215,13"),
    ]
    y = 800
    for x, texte in haut:
        cv.drawString(x, y, texte)
        y -= 16
    cv.setFont("Helvetica", 5)   # petite police, comme sur une vraie fiche : titres de colonnes serrés
    y = 200
    for x, titre in [(40, "Cumuls"), (75, "Hrs travaillées"), (121, "Salaire brut"), (165, "Brut abattu"),
                     (252, "Tranche 1"), (295, "Tranche 2"), (332, "Net imposable"), (376, "Impôt prélevé"),
                     (417, "HS/HC exonér"), (456, "Allèg. cot. employeur")]:
        cv.drawString(x, y, titre)
    colonnes_droite = [110, 153, 196, 282, 325, 368, 410, 451]
    lignes = [("Mensuel", ["22,20", "273,28", "273,28", "273,28", "0,00", net[0], impot[0], hs[0]]),
              ("Annuel", ["22,20", "1 273,28", "1 273,28", "1 273,28", "0,00", net[1], impot[1], hs[1]])]
    for nom, valeurs in lignes:
        y -= 12
        cv.drawString(40, y, nom)
        for xd, val in zip(colonnes_droite, valeurs):
            cv.drawRightString(xd, y, val)
    cv.save()


def test_tableau_cumuls(tmp_path):
    pytest.importorskip("reportlab")
    pdf = tmp_path / "fiche.pdf"
    _pdf_tableau_cumuls(pdf)
    assert c.extraire_tableau_cumuls(pdf) == {
        "net_imposable": {"mois": 222.95, "cumul": 1456.78},
        "impot_preleve": {"mois": 18.5, "cumul": 18.5},
        "hs_exonerees": {"mois": 12.0, "cumul": 12.0},
    }


def test_analyse_fiche_avec_tableau_cumuls(tmp_path):
    pytest.importorskip("reportlab")
    pdf = tmp_path / "fiche.pdf"
    _pdf_tableau_cumuls(pdf)
    r = c.analyser(pdf)
    assert (r["annee"], r["mois"]) == (2026, 9)
    assert r["employeur"] == "Job & Box"
    assert r["net_imposable"] == 222.95
    assert r["cumul"] == 1456.78
    assert r["net_a_payer"] == 215.13
    assert r["brut"] == 273.28
    assert r["impot_preleve"] == 18.5
    assert r["hs_exonerees"] == 12.0
    assert r["avert"] == []


def test_report_impot_preleve_et_hs(tmp_path):
    """Deux bulletins du même employeur : les montants s'additionnent et sont contrôlés
    avec le cumul annuel de la dernière fiche."""
    pytest.importorskip("reportlab")
    src = tmp_path / "pdf"
    src.mkdir()
    _pdf_tableau_cumuls(src / "sept.pdf", "01/09/2026", "30/09/2026",
                        net=("1 000,00", "1 000,00"), impot=("18,50", "18,50"), hs=("12,00", "12,00"))
    _pdf_tableau_cumuls(src / "oct.pdf", "01/10/2026", "31/10/2026",
                        net=("1 100,00", "2 100,00"), impot=("20,00", "38,50"), hs=("0,00", "12,00"))
    res = c.traiter(src, tmp_path / "out", log=lambda m: None)
    par_case = {(l["annee"], l["case"]): l for l in res["lignes"]}
    assert par_case[(2026, "1AJ")]["total"] == pytest.approx(2100.0)
    assert par_case[(2026, "8HV")]["total"] == pytest.approx(38.5)
    assert par_case[(2026, "1GH")]["total"] == pytest.approx(12.0)
    for cle in ((2026, "1AJ"), (2026, "8HV"), (2026, "1GH")):
        assert par_case[cle]["raisons"] == []
    page = res["html"].read_text(encoding="utf-8")
    assert "8HV" in page and "1GH" in page


def test_impot_prelevé_incoherent_avec_le_cumul(tmp_path):
    """Si la somme des fiches ne colle pas avec le cumul annuel, l'année est à vérifier."""
    pytest.importorskip("reportlab")
    src = tmp_path / "pdf"
    src.mkdir()
    _pdf_tableau_cumuls(src / "oct.pdf", "01/10/2026", "31/10/2026",
                        net=("1 100,00", "1 100,00"), impot=("20,00", "58,50"), hs=("0,00", "0,00"))
    res = c.traiter(src, tmp_path / "out", log=lambda m: None)
    par_case = {(l["annee"], l["case"]): l for l in res["lignes"]}
    assert par_case[(2026, "8HV")]["raisons"]          # écart avec le cumul signalé
    assert (2026, "1GH") not in par_case                # aucune HS : pas de ligne inutile


def test_colonnes_absentes_donnent_non_lu(tmp_path):
    """Un bulletin sans tableau « Cumuls » : 8HV est signalé « non lu », pas à 0 en silence."""
    pytest.importorskip("reportlab")
    src = tmp_path / "pdf"
    src.mkdir()
    _pdf(src / "b1.pdf", ["Employeur : SARL Dupont Informatique", "Bulletin de salaire - mois de octobre 2025",
                          "Total brut 2 100,50", "Net fiscal 1 650,25"])
    res = c.traiter(src, tmp_path / "out", log=lambda m: None)
    par_case = {(l["annee"], l["case"]): l for l in res["lignes"]}
    assert any("non lu" in r for r in par_case[(2025, "8HV")]["raisons"])


def test_declarant_invalide(tmp_path):
    with pytest.raises(ValueError):
        c.traiter(tmp_path, tmp_path / "out", log=lambda m: None, declarant="9")


def test_dossier_sans_pdf(tmp_path):
    with pytest.raises(ValueError):
        c.traiter(tmp_path, tmp_path / "out", log=lambda m: None)


def test_dossier_introuvable(tmp_path):
    with pytest.raises(ValueError):
        c.traiter(tmp_path / "nexiste_pas", tmp_path / "out", log=lambda m: None)
